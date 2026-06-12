"""Exportación de cronogramas a Excel (.xlsx) con openpyxl."""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..core.modelos import CuotaRegistro, Prestamo

ENCABEZADOS = ["N°", "Fecha", "Días", "Amortización", "Interés", "Cuota", "Saldo Pendiente"]

_AZUL = "1F4E78"
_GRIS = "F2F2F2"


def _f(valor: Decimal) -> float:
    return float(valor)


def exportar_cronograma(
    prestamo: Prestamo,
    ruta: str | Path,
    cuotas: list[CuotaRegistro] | None = None,
    titulo: str | None = None,
) -> Path:
    """Escribe el cronograma del préstamo (o uno provisto, p. ej. una
    simulación) en un archivo .xlsx con el formato de columnas estándar."""
    cuotas = cuotas if cuotas is not None else prestamo.cuotas
    ruta = Path(ruta)

    wb = Workbook()
    ws = wb.active
    ws.title = "Cronograma"

    borde = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
    moneda = '#,##0.00'

    # Encabezado de información del préstamo.
    titulo = titulo or f"Cronograma - {prestamo.cliente.nombre}"
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=14, color=_AZUL)
    ws["A2"] = (
        f"Monto: S/ {_f(prestamo.monto):,.2f}    "
        f"Tasa mensual: {_f(prestamo.tasa_mensual) * 100:.4f}%    "
        f"Fórmula: {prestamo.formula}    "
        f"Desembolso: {prestamo.fecha_desembolso.isoformat()}    "
        f"1er venc.: {prestamo.fecha_primer_vencimiento.isoformat()}    "
        f"Cuotas: {len(cuotas)}"
    )
    ws["A2"].font = Font(italic=True, color="595959")

    # Fila de encabezados de la tabla.
    fila0 = 4
    for col, texto in enumerate(ENCABEZADOS, start=1):
        celda = ws.cell(row=fila0, column=col, value=texto)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=_AZUL)
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = borde

    # Filas de datos.
    for i, c in enumerate(cuotas):
        fila = fila0 + 1 + i
        valores = [
            c.numero, c.fecha.isoformat(), c.dias,
            _f(c.amortizacion), _f(c.interes), _f(c.cuota), _f(c.saldo),
        ]
        for col, valor in enumerate(valores, start=1):
            celda = ws.cell(row=fila, column=col, value=valor)
            celda.border = borde
            if col >= 4:
                celda.number_format = moneda
            if col in (1, 3):
                celda.alignment = Alignment(horizontal="center")
        if i % 2 == 1:
            for col in range(1, len(ENCABEZADOS) + 1):
                ws.cell(row=fila, column=col).fill = PatternFill("solid", fgColor=_GRIS)

    # Fila de totales.
    fila_total = fila0 + 1 + len(cuotas)
    ws.cell(row=fila_total, column=3, value="TOTAL").font = Font(bold=True)
    total_cuota = sum(_f(c.cuota) for c in cuotas)
    total_interes = sum(_f(c.interes) for c in cuotas)
    total_amort = sum(_f(c.amortizacion) for c in cuotas)
    for col, valor in ((4, total_amort), (5, total_interes), (6, total_cuota)):
        celda = ws.cell(row=fila_total, column=col, value=valor)
        celda.font = Font(bold=True)
        celda.number_format = moneda

    # Anchos de columna.
    anchos = [6, 14, 7, 16, 14, 14, 18]
    for col, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho

    wb.save(ruta)
    return ruta
